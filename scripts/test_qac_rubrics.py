"""Phase 1 integrity checks for the QAC rubric library and Excel templates.

Run from the repository root:
    python -B -m unittest discover -s scripts -p test_qac_rubrics.py

Optional: set QAC_ORIGINALS_DIR to the folder holding the four workbooks as supplied by QAC to also
compare the stored templates byte-for-byte against them (the manifest SHA-256 values were taken from
those files). openpyxl, when installed, is used as an independent second reader of the workbooks.
"""
import hashlib
import json
import os
import re
import re
import shutil
import subprocess
import tempfile
import unittest
import warnings
import zipfile
from pathlib import Path

import generate_qac_rubrics as gen
from qac_xlsx import Workbook, num_to_col, sha256_file, split_cell, split_range

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / 'data'
BASELINE_COMMIT = '1861291'  # last commit before the Phase 1 foundation work
PHASE1_MAY_MODIFY = {'.gitattributes', 'data/README.md'}
EVIDENCE_PATH = 'data/ee_alignment_evidence.json'
CURRICULUM_PATH = 'data/ee_curriculum.json'
# Owner decision D4: the only curriculum / guidance edits allowed after the baseline.
PI31_ORIGINAL = 'Communicate relevant, accurate, and sufficient engineering information.'
PI31_ADOPTED = 'Communicate engineering information using appropriate technical content.'
DEFAULT_ORIGINALS = Path('C:/Users/user/Downloads/New folder')
ORIGINAL_NAMES = [t['file'] for t in gen.TEMPLATES]


def load(name):
    return json.loads((DATA / name).read_text(encoding='utf-8'))


def git(*args):
    return subprocess.run(['git', *args], cwd=ROOT, capture_output=True, text=True, encoding='utf-8')


RUBRICS = load('ee-qac-rubrics.json')
TEMPLATES = load('ee-qac-templates.json')
MAPPING = load('ee-qac-main-sheet-mapping.json')
ISSUES = load('ee-qac-review-issues.json')
CURRICULUM = load('ee_curriculum.json')
SPEC = {t['template_id']: t for t in gen.TEMPLATES}


def open_template(template_id):
    return Workbook(gen.template_path(SPEC[template_id]))


def t_sha(template_id):
    return next(t['sha256'] for t in TEMPLATES['templates'] if t['template_id'] == template_id)


def clear_main_row(source, target, row):
    """Test-only model of the future exporter step: empty every cell of one Main row in a COPY.

    Cell value, formula and cached value are removed; the style attribute is kept. The Main
    sheet part and (when present) the matching calcChain entries are the only parts rewritten.
    """
    with Workbook(source) as wb:
        main = wb.sheet('Main')
        part = main.path
        workbook_xml = wb.zip.read('xl/workbook.xml').decode('utf-8')
        sheet_id = re.search(r'<sheet [^>]*name="Main"[^>]*sheetId="(\d+)"|<sheet [^>]*sheetId="(\d+)"[^>]*name="Main"', workbook_xml)
        sheet_id = next(g for g in sheet_id.groups() if g)
    with zipfile.ZipFile(source) as zin, zipfile.ZipFile(target, 'w', zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == part:
                text = data.decode('utf-8')
                def clear_row(m):
                    def clear_cell(c):
                        ref = re.search(r'\br="([A-Z]+%d)"' % row, c.group(0))
                        style = re.search(r'\bs="(\d+)"', c.group(0).split('>')[0])
                        return '<c r="%s"%s/>' % (ref.group(1), ' s="%s"' % style.group(1) if style else '')
                    return re.sub(r'<c\b[^>]*?/>|<c\b[^>]*?>.*?</c>', clear_cell, m.group(0), flags=re.S)
                new_text, n = re.subn(r'<row\b[^>]*\br="%d"[^>]*>.*?</row>' % row, clear_row, text, flags=re.S)
                assert n == 1
                data = new_text.encode('utf-8')
            elif info.filename == 'xl/calcChain.xml':
                text = data.decode('utf-8')
                text = re.sub(r'<c\b[^>]*\br="[A-Z]+%d"[^>]*\bi="%s"[^>]*/>|<c\b[^>]*\bi="%s"[^>]*\br="[A-Z]+%d"[^>]*/>' % (row, sheet_id, sheet_id, row), '', text)
                data = text.encode('utf-8')
            zout.writestr(info, data)


class TemplateAssetTests(unittest.TestCase):
    def test_four_templates_listed_and_present(self):
        self.assertEqual([t['so'] for t in TEMPLATES['templates']], ['SO1', 'SO2', 'SO3', 'SO4'])
        folder_files = sorted(p.relative_to(ROOT).as_posix() for p in (ROOT / TEMPLATES['asset_root']).rglob('*') if p.is_file())
        self.assertEqual(folder_files, sorted(t['asset_path'] for t in TEMPLATES['templates']),
                         'template folder must contain exactly the four stored workbooks')

    def test_checksums_match_manifest(self):
        for t in TEMPLATES['templates']:
            path = ROOT / t['asset_path']
            with self.subTest(template=t['template_id']):
                self.assertEqual(path.name, t['original_file_name'])
                self.assertEqual(path.stat().st_size, t['size_bytes'])
                self.assertEqual(sha256_file(path), t['sha256'])
                self.assertIsNone(zipfile.ZipFile(path).testzip())

    def test_git_does_not_normalise_templates(self):
        result = git('check-attr', 'text', 'diff', '--', TEMPLATES['templates'][0]['asset_path'])
        if result.returncode != 0:
            self.skipTest('git unavailable')
        self.assertIn('text: unset', result.stdout)

    def test_identical_to_supplied_originals(self):
        folder = Path(os.environ.get('QAC_ORIGINALS_DIR', DEFAULT_ORIGINALS))
        if not folder.is_dir():
            self.skipTest('Supplied originals folder not available (set QAC_ORIGINALS_DIR)')
        for t in TEMPLATES['templates']:
            original = folder / t['original_file_name']
            with self.subTest(template=t['template_id']):
                self.assertTrue(original.is_file(), original)
                self.assertEqual((ROOT / t['asset_path']).read_bytes(), original.read_bytes())

    def test_sheet_inventory_matches_workbooks(self):
        for t in TEMPLATES['templates']:
            with self.subTest(template=t['template_id']), open_template(t['template_id']) as wb:
                self.assertEqual(wb.sheet_names, t['sheet_order'])
                self.assertEqual([(s['name'], s['state']) for s in t['sheets']], [(s.name, s.state) for s in wb.sheets])
                self.assertEqual(wb.sheet_names[0], 'Main')
                for s in t['sheets']:
                    if s['role'] == 'authoritative_rubric':
                        self.assertEqual(s['state'], 'visible')
                    if s['role'] == 'legacy_hidden':
                        self.assertEqual(s['state'], 'hidden')
                self.assertEqual(sorted(s['name'] for s in t['sheets'] if s['role'] == 'authoritative_rubric'),
                                 sorted(t['rubric_sheets']))


class RubricJsonTests(unittest.TestCase):
    def test_structure(self):
        self.assertEqual(RUBRICS['schema_version'], '1.0')
        ids = [r['rubric_id'] for r in RUBRICS['rubrics']]
        self.assertEqual(len(ids), len(set(ids)))
        self.assertEqual(sorted({r['so'] for r in RUBRICS['rubrics']}), ['SO1', 'SO2', 'SO3', 'SO4'])
        criterion_ids = [c['criterion_id'] for r in RUBRICS['rubrics'] for c in r['criteria']]
        self.assertEqual(len(criterion_ids), len(set(criterion_ids)))
        self.assertEqual(len(criterion_ids), 3 + 8 + 3 + 3 + 2)
        known_templates = {t['template_id'] for t in TEMPLATES['templates']}
        for r in RUBRICS['rubrics']:
            with self.subTest(rubric=r['rubric_id']):
                self.assertIn(r['template_id'], known_templates)
                self.assertIsNone(r['weights'])
                self.assertEqual([x['score'] for x in r['performance_levels']['scale']], [3, 2, 1, 0])
                for field in ('workbook', 'sheet', 'sheet_state', 'workbook_version_label'):
                    self.assertTrue(r['source'][field])
                self.assertEqual(r['source']['sheet_state'], 'visible')

    def test_every_criterion_complete(self):
        for r in RUBRICS['rubrics']:
            for c in r['criteria']:
                with self.subTest(criterion=c['criterion_id']):
                    self.assertEqual(list(c['descriptors']), ['above', 'meet', 'developing', 'below'])
                    for level, d in c['descriptors'].items():
                        self.assertTrue(d['text'].strip(), level)
                    self.assertEqual([d['score'] for d in c['descriptors'].values()], [3, 2, 1, 0])
                    self.assertTrue(c['label'].strip())
                    self.assertGreaterEqual(len(c['checklists']), 1)
                    self.assertEqual(len(c['checklists']), len(r['checklist_columns']))
                    for cl in c['checklists']:
                        self.assertEqual(''.join(s['text'] for s in cl['segments']), cl['text'])
                        self.assertIn(cl['role'], ('project', 'exam', 'unspecified'))

    def test_team_and_individual_distinction_preserved(self):
        for r in RUBRICS['rubrics']:
            for c in r['criteria']:
                for cl in c['checklists']:
                    audiences = [s['audience'] for s in cl['segments']]
                    with self.subTest(criterion=c['criterion_id'], checklist=cl['checklist_id']):
                        if cl['role'] == 'exam':
                            self.assertEqual(audiences, ['unspecified'])
                            self.assertEqual(cl['segments'][0]['context'], 'exam_written')
                        else:
                            self.assertEqual(audiences, ['team', 'individual'])

    def test_no_non_qac_content(self):
        forbidden = re.compile(r'(suggest|assignment|clo_|instruction|guidance)', re.I)
        def walk(node, path=''):
            if isinstance(node, dict):
                for k, v in node.items():
                    self.assertIsNone(forbidden.search(k), 'key %s%s' % (path, k))
                    walk(v, path + k + '.')
            elif isinstance(node, list):
                for v in node:
                    walk(v, path)
        walk([RUBRICS['rubrics'], RUBRICS['index']])

    def test_generated_files_are_up_to_date(self):
        expectations = {
            'ee-qac-rubrics.json': gen.build_rubric_library(CURRICULUM),
            'ee-qac-templates.json': gen.build_template_manifest(),
        }
        import qac_mapping_data
        expectations['ee-qac-main-sheet-mapping.json'] = qac_mapping_data.build_mapping(CURRICULUM)
        expectations['ee-qac-review-issues.json'] = qac_mapping_data.build_issues()
        for name, value in expectations.items():
            with self.subTest(file=name):
                self.assertEqual((DATA / name).read_text(encoding='utf-8'),
                                 json.dumps(value, indent=2, ensure_ascii=False) + '\n')


class SourceFidelityTests(unittest.TestCase):
    def rubric_cells(self, r):
        cells = {r['title']['cell'], r['item_header']['cell'], r['assessment_note']['cell']}
        cells |= {h['cell'] for h in r['performance_levels']['headers'].values()}
        cells |= {c['header_cell'] for c in r['checklist_columns']}
        for c in r['criteria']:
            cells.add(c['label_cell'])
            cells |= {d['cell'] for d in c['descriptors'].values()}
            cells |= {cl['cell'] for cl in c['checklists']}
        return cells

    def test_every_text_equals_source_cell(self):
        for r in RUBRICS['rubrics']:
            with self.subTest(rubric=r['rubric_id']), open_template(r['template_id']) as wb:
                sheet = wb.sheet(r['source']['sheet'])
                self.assertEqual(sheet.state, r['source']['sheet_state'])
                self.assertEqual(sheet.dimension, r['source']['dimension'])
                self.assertEqual(sheet.merged, r['source']['merged_ranges'])
                self.assertEqual(sheet.value(r['title']['cell']), r['title']['text'])
                self.assertEqual(sheet.value(r['item_header']['cell']), r['item_header']['text'])
                self.assertEqual(sheet.value(r['assessment_note']['cell']), r['assessment_note']['text'])
                for h in r['performance_levels']['headers'].values():
                    self.assertEqual(sheet.value(h['cell']), h['text'])
                for c in r['checklist_columns']:
                    self.assertEqual(sheet.value(c['header_cell']), c['header'])
                for c in r['criteria']:
                    self.assertEqual(sheet.value(c['label_cell']), c['label'])
                    for d in c['descriptors'].values():
                        self.assertEqual(sheet.value(d['cell']), d['text'])
                    for cl in c['checklists']:
                        self.assertEqual(sheet.value(cl['cell']), cl['text'])

    def test_no_rubric_cell_left_out(self):
        for r in RUBRICS['rubrics']:
            with self.subTest(rubric=r['rubric_id']), open_template(r['template_id']) as wb:
                sheet = wb.sheet(r['source']['sheet'])
                self.assertEqual(set(sheet.cells), self.rubric_cells(r),
                                 'extracted cells must equal the non-empty cells of the rubric sheet')

    def test_agrees_with_independent_reader(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest('openpyxl not installed')
        warnings.simplefilter('ignore')
        for r in RUBRICS['rubrics']:
            with self.subTest(rubric=r['rubric_id']):
                ws = openpyxl.load_workbook(gen.template_path(SPEC[r['template_id']]))[r['source']['sheet']]
                self.assertEqual(ws.sheet_state, r['source']['sheet_state'])
                for ref in self.rubric_cells(r):
                    self.assertIsNotNone(ws[ref].value, ref)
                for c in r['criteria']:
                    self.assertEqual(ws[c['label_cell']].value, c['label'])
                    for d in c['descriptors'].values():
                        self.assertEqual(ws[d['cell']].value, d['text'])
                    for cl in c['checklists']:
                        self.assertEqual(ws[cl['cell']].value, cl['text'])

    def test_scoring_rules_match_workbook_formulas(self):
        for rule in RUBRICS['scoring_rules']:
            with self.subTest(template=rule['template_id']), open_template(rule['template_id']) as wb:
                sheet = wb.sheet(rule['source']['sheet'])
                self.assertEqual(sheet.protected, rule['source']['sheet_protected'])
                self.assertIsNone(rule['weights'])
                self.assertEqual(sheet.value(rule['notes']['cell']), rule['notes']['text'])
                self.assertEqual(sheet.value(rule['score_entry_instruction']['cell']), rule['score_entry_instruction']['text'])
                def check(entry):
                    self.assertEqual(sheet.formula(entry['cell']), entry['formula'])
                    self.assertTrue(entry['formula'])
                for key, entry in rule['formulas'].items():
                    for item in (entry if isinstance(entry, list) else [entry]):
                        check(item)
                att = rule['formulas']['student_attainment']['formula']
                m = re.search(r'<1,"BE",IF\(\w+<2,"DE",IF\(\w+<3,"ME","AE"\)\)\)', att)
                self.assertTrue(m, att)
                bands = {b['code']: (b['min_inclusive'], b['max_exclusive']) for b in rule['attainment_thresholds']['bands']}
                self.assertEqual(bands, {'BE': (0, 1), 'DE': (1, 2), 'ME': (2, 3), 'AE': (3, None)})


class CurriculumAssociationTests(unittest.TestCase):
    def test_so_pi_associations(self):
        abet = CURRICULUM['abet']
        for r in RUBRICS['rubrics']:
            self.assertIn(r['so'], abet['student_outcomes'])
            for c in r['criteria']:
                with self.subTest(criterion=c['criterion_id']):
                    self.assertEqual(c['so'], r['so'])
                    self.assertIn(c['pi'], abet['performance_indicators'])
                    self.assertEqual(abet['performance_indicators'][c['pi']]['so'], r['so'])
                    self.assertIn(c['pi'], abet['student_outcomes'][r['so']]['pis'])
                    if c['cip_pi_statement_comparison'] != 'not_applicable':
                        self.assertEqual(c['cip_pi_statement_comparison'],
                                         gen.compare_with_cip(c['label'], abet['performance_indicators'][c['pi']]['statement']))
                        self.assertEqual(gen.pi_id(c['label']), c['pi'])

    def test_each_so_rubrics_cover_every_cip_pi_of_that_so(self):
        for so, info in CURRICULUM['abet']['student_outcomes'].items():
            covered = sorted({p for r in RUBRICS['rubrics'] if r['so'] == so for p in r['pis']})
            if so in ('SO1', 'SO2', 'SO3', 'SO4'):
                with self.subTest(so=so):
                    self.assertEqual(covered, sorted(info['pis']))
            else:
                self.assertEqual(covered, [])

    def test_rubric_pis_match_main_sheet_pi_list(self):
        for spec in gen.TEMPLATES:
            with self.subTest(template=spec['template_id']), open_template(spec['template_id']) as wb:
                text = wb.sheet('Main').value('F5')
                listed = sorted(set(re.findall(r'PI\s?(\d\d)', text)))
                rubric_pis = sorted({p[2:] for r in RUBRICS['rubrics'] if r['template_id'] == spec['template_id'] for p in r['pis']})
                self.assertEqual(listed, rubric_pis)

    def test_rubric_json_does_not_copy_cip_descriptions(self):
        statements = [i['statement'] for i in CURRICULUM['abet']['student_outcomes'].values()]
        blob = json.dumps(RUBRICS, ensure_ascii=False)
        for s in statements:
            self.assertNotIn(s, blob)


class Pi31AuthoritativeWordingTests(unittest.TestCase):
    def test_cip_pi31_equals_visible_qac_rubric_wording(self):
        pi31 = CURRICULUM['abet']['performance_indicators']['PI31']
        self.assertEqual(pi31['statement'], PI31_ADOPTED)
        self.assertEqual((pi31['so'], pi31['rubric']), ('SO3', 'Rubric'))
        self.assertEqual(CURRICULUM['abet']['student_outcomes']['SO3']['pis'], ['PI31', 'PI32', 'PI33'])
        criterion = next(c for r in RUBRICS['rubrics'] for c in r['criteria'] if c['pi'] == 'PI31')
        self.assertEqual(gen.pi_label_text(criterion['label']) + '.', pi31['statement'])
        self.assertEqual(criterion['cip_pi_statement_comparison'], 'punctuation_or_case_only')
        with open_template('qac-so3-ver01') as wb:
            self.assertEqual(wb.sheet('Rubric').value('A4').strip(), 'PI31 – ' + pi31['statement'][:-1])
            self.assertEqual(wb.sheet('Main').value('F5').split('\n')[0], 'PI31 – ' + pi31['statement'])

    def test_pi31_mappings_preserved(self):
        for course in CURRICULUM['curriculum']['courses']:
            for clo in course['clos']:
                if 'PI31' in clo['pi_codes']:
                    self.assertIn('SO3', clo['mapped_sos'], course['course_code'])

    def test_alignment_evidence_profile_reflects_adopted_wording(self):
        profile = load('ee_alignment_evidence.json')['pi_evidence_profiles']['PI31']
        blob = json.dumps(profile)
        self.assertNotIn('relevant, accurate', blob)
        self.assertIn('appropriate technical content', profile['demonstrates'])
        self.assertIn('appropriate technical content', ' '.join(profile['rubric_criteria']))
        self.assertEqual(profile['primary_family'], 'communication')

    def test_wording_recorded_in_decisions(self):
        d4 = MAPPING['decisions']['D4']
        self.assertEqual((d4['original'], d4['replacement']), (PI31_ORIGINAL, PI31_ADOPTED))


class DecisionRequirementTests(unittest.TestCase):
    def test_assessment_limit_recorded(self):
        limits = MAPPING['preparation_limits']
        self.assertEqual(limits['max_assessment_activities_per_session'], 5)
        self.assertIn('total', limits['scope'])
        for t in MAPPING['templates']:
            self.assertEqual(t['assessment_block']['max_activities_per_session'], 5)

    def test_one_so_per_workbook_and_row6_cleared_by_rule(self):
        for t in MAPPING['templates']:
            with self.subTest(template=t['template_id']), open_template(t['template_id']) as wb:
                main = wb.sheet('Main')
                self.assertEqual([r['row'] for r in t['so_rows']], [5])
                self.assertEqual(main.value('A5'), t['so'], 'row 5 must hold the workbook SO')
                self.assertEqual(len(t['cleared_rows']), 1)
                rule = t['cleared_rows'][0]
                self.assertEqual((rule['row'], rule['action'], rule['decision']), (6, 'clear_row', 'D2'))
                self.assertEqual([c['cell'] for c in rule['cells_with_template_content']], ['A6', 'B6', 'F6', 'J6'])
                self.assertEqual(main.value('A6'), 'SO4')
                self.assertTrue(main.value('J6').startswith('3.1 Recognize the engineering profession'))
                self.assertIn('VLOOKUP(A6', main.formula('B6'))
                for c in rule['cells_with_template_content']:
                    self.assertEqual(c['template_content']['kind'], 'formula' if c['cell'] == 'B6' else 'prefilled')
                self.assertEqual(rule['preserve'][0], 'merged ranges')

    def test_generated_copy_row6_is_cleared_and_everything_else_preserved(self):
        for t in MAPPING['templates']:
            with self.subTest(template=t['template_id']), tempfile.TemporaryDirectory() as tmp:
                source = gen.template_path(SPEC[t['template_id']])
                before = sha256_file(source)
                copy = Path(tmp) / 'copy.xlsx'
                clear_main_row(source, copy, 6)
                with Workbook(source) as orig, Workbook(copy) as new:
                    om, nm = orig.sheet('Main'), new.sheet('Main')
                    for ref in ('A6', 'B6', 'F6', 'J6'):
                        self.assertIsNotNone(om.cells.get(ref))
                        self.assertIsNone(nm.cells.get(ref), ref + ' must be empty')
                    self.assertFalse([r for r in nm.cells if r.endswith('6') and r[:-1].isalpha() and r[-1] == '6' and r[:-1] in 'ABCDEFGHIJ'])
                    self.assertEqual({k: v for k, v in om.cells.items() if split_cell(k)[1] != 6},
                                     {k: v for k, v in nm.cells.items() if split_cell(k)[1] != 6})
                    self.assertEqual(om.merged, nm.merged)
                    self.assertEqual(om.data_validations, nm.data_validations)
                    self.assertEqual(om.comments, nm.comments)
                    self.assertEqual(orig.sheet_names, new.sheet_names)
                    self.assertEqual([(s.name, s.state) for s in orig.sheets], [(s.name, s.state) for s in new.sheets])
                    for a, b in zip(orig.sheets, new.sheets):
                        if a.name != 'Main':
                            self.assertEqual(a.cells, b.cells)
                    self.assertFalse(nm.formulas_in('A6:J6'))
                # byte-level: only the Main sheet part and calcChain may differ
                with zipfile.ZipFile(source) as zo, zipfile.ZipFile(copy) as zn:
                    self.assertIsNone(zn.testzip())
                    self.assertEqual(zo.namelist(), zn.namelist())
                    main_part = Workbook(source).sheet('Main').path
                    differing = {n for n in zo.namelist() if zo.read(n) != zn.read(n)}
                    self.assertTrue(differing <= {main_part, 'xl/calcChain.xml'}, differing)
                    if 'xl/calcChain.xml' in zo.namelist():
                        main_id = re.search(r'name="Main"[^>]*sheetId="(\d+)"|sheetId="(\d+)"[^>]*name="Main"', zo.read('xl/workbook.xml').decode()).groups()
                        main_id = next(g for g in main_id if g)
                        self.assertIsNone(re.search(r'<c [^>]*r="[A-Z]+6"[^>]*i="%s"' % main_id, zn.read('xl/calcChain.xml').decode()))
                self.assertEqual(sha256_file(source), before, 'original template must not change')
                self.assertEqual(sha256_file(source), t_sha(t['template_id']))

    def test_mapped_pi_rule_is_selection_based(self):
        rule = ' '.join(MAPPING['decisions']['D3']['derivation'])
        self.assertIn('instructor-selected', rule)
        self.assertIn('without duplicates', rule)
        self.assertIn('traceability', MAPPING['decisions']['D3']['derivation'][-1])
        for t in MAPPING['templates']:
            for row in t['so_rows']:
                f = next(x for x in row['fields'] if x['field_id'].startswith('mapped_pis'))
                self.assertIn('D3', f['requirement_rules'])
                self.assertNotIn('all PIs', f['cip_source']['path'].replace('PIs in', ''))

    def test_issue_classification_follows_decisions(self):
        by_id = {i['id']: i for i in ISSUES['issues']}
        self.assertEqual(by_id['QAC-10']['status'], 'resolved_by_decision')
        self.assertEqual(by_id['QAC-11']['status'], 'resolved_by_decision')
        self.assertEqual(by_id['QAC-13']['status'], 'partially_resolved')
        self.assertEqual(by_id['QAC-15']['status'], 'partially_resolved')
        self.assertEqual(by_id['QAC-20']['status'], 'resolved_by_decision')
        self.assertEqual(by_id['QAC-21']['status'], 'resolved_by_decision')
        self.assertEqual(by_id['QAC-10']['resolution']['decision'], 'D2')


class RetrievalTests(unittest.TestCase):
    def setUp(self):
        self.by_id = {r['rubric_id']: r for r in RUBRICS['rubrics']}
        self.crit = {c['criterion_id']: (r, c) for r in RUBRICS['rubrics'] for c in r['criteria']}

    def criteria_for_pi(self, pi):
        return [self.crit[cid][1] for cid in RUBRICS['index']['by_pi'].get(pi, [])]

    def test_index_matches_brute_force_scan(self):
        by_so, by_pi, by_context = {}, {}, {}
        for r in RUBRICS['rubrics']:
            by_so.setdefault(r['so'], []).append(r['rubric_id'])
            for c in r['criteria']:
                by_pi.setdefault(c['pi'], []).append(c['criterion_id'])
                for cl in c['checklists']:
                    for s in cl['segments']:
                        by_context.setdefault(s['context'], set()).add(c['criterion_id'])
        index = RUBRICS['index']
        self.assertEqual(index['by_so'], by_so)
        self.assertEqual(index['by_pi'], by_pi)
        self.assertEqual({k: set(v) for k, v in index['by_checklist_context'].items()}, by_context)
        self.assertEqual({cid: r['rubric_id'] for cid, (r, c) in self.crit.items()}, index['by_criterion'])

    def test_retrieve_by_so_and_pi(self):
        expected = {'PI11': 1, 'PI12': 1, 'PI13': 1, 'PI21': 8, 'PI22': 1, 'PI23': 1, 'PI24': 1,
                    'PI31': 1, 'PI32': 1, 'PI33': 1, 'PI41': 1, 'PI42': 1}
        for pi, count in expected.items():
            with self.subTest(pi=pi):
                found = self.criteria_for_pi(pi)
                self.assertEqual(len(found), count)
                self.assertTrue(all(c['pi'] == pi for c in found))
        self.assertEqual(len(RUBRICS['index']['by_so']['SO2']), 2)
        self.assertEqual(self.criteria_for_pi('PI55'), [])

    def test_retrieve_by_criterion_and_context(self):
        r, c = self.crit['PI21.C3']
        self.assertEqual(r['so'], 'SO2')
        self.assertTrue(c['label'].startswith('3. Alternatives'))
        individual = [s for cl in c['checklists'] for s in cl['segments'] if s['context'] == 'individual_verification']
        self.assertEqual(len(individual), 1)
        self.assertIn('Identify one alternative', individual[0]['text'])
        exam = RUBRICS['index']['by_checklist_context']['exam_written']
        self.assertEqual(sorted(exam), ['PI11.C1', 'PI12.C1', 'PI13.C1', 'PI41.C1', 'PI42.C1'])


class MainSheetMappingTests(unittest.TestCase):
    def all_fields(self, t):
        fields = list(t['header_fields'])
        for row in t['so_rows']:
            fields += row['fields']
        fields += t['assessment_block']['fields'] + t['template_metadata_fields']
        return fields

    def test_one_mapping_per_template(self):
        self.assertEqual([t['template_id'] for t in MAPPING['templates']], [t['template_id'] for t in TEMPLATES['templates']])

    def test_targets_are_valid_cells_and_ranges(self):
        issue_ids = {i['id'] for i in ISSUES['issues']}
        for t in MAPPING['templates']:
            with self.subTest(template=t['template_id']), open_template(t['template_id']) as wb:
                main = wb.sheet(t['main_sheet'])
                max_col, max_row = split_cell(main.dimension.split(':')[1])
                seen = set()
                for f in self.all_fields(t):
                    self.assertNotIn(f['field_id'], seen)
                    seen.add(f['field_id'])
                    tgt = f['target']
                    c1, r1, c2, r2 = split_range(tgt['range'])
                    self.assertTrue(1 <= c1 <= c2 <= max_col and 1 <= r1 <= r2 <= max_row, f['field_id'])
                    merged = main.merged_range_for(tgt['anchor'])
                    self.assertEqual(tgt['merged'], bool(merged))
                    self.assertEqual(tgt['range'], merged or tgt['anchor'])
                    self.assertEqual(split_cell(tgt['anchor']), (c1, r1), 'anchor must be the top-left cell')
                    if f['label']:
                        self.assertEqual(main.value(f['label']['cell']), f['label']['text'])
                        self.assertTrue(f['label']['text'])
                    for flag in f['flags']:
                        self.assertIn(flag, issue_ids)
                    self.assertIn(f['entry_mode'], MAPPING['vocabulary']['entry_mode'])
                    self.assertIn(f['requirement'], MAPPING['vocabulary']['requirement'])

    def test_current_content_matches_template(self):
        for t in MAPPING['templates']:
            with self.subTest(template=t['template_id']), open_template(t['template_id']) as wb:
                main = wb.sheet('Main')
                for f in self.all_fields(t):
                    anchor, cc = f['target']['anchor'], f['current_content']
                    value, formula = main.value(anchor), main.formula(anchor)
                    if cc['kind'] == 'empty':
                        self.assertIsNone(value, f['field_id'])
                        self.assertIsNone(formula, f['field_id'])
                    elif cc['kind'] == 'formula':
                        self.assertEqual(formula, cc['formula'])
                        self.assertEqual(f['entry_mode'], 'workbook_formula')
                        self.assertFalse(f['editable'])
                    else:
                        text = value if isinstance(value, str) else str(value)
                        self.assertEqual(hashlib.sha256(text.encode('utf-8')).hexdigest(), cc['text_sha256'], f['field_id'])

    def test_selector_validations_exist(self):
        for t in MAPPING['templates']:
            with self.subTest(template=t['template_id']), open_template(t['template_id']) as wb:
                main = wb.sheet('Main')
                validations = main.data_validations
                self.assertTrue(validations)
                for row in t['so_rows']:
                    f = row['fields'][0]
                    self.assertEqual(f['entry_mode'], 'selector')
                    dv = f['data_validation']
                    self.assertIsNotNone(dv)
                    self.assertTrue(any(v['formula1'] == dv['source'] and f['target']['anchor'] in (v['sqref'] or '') or
                                        (v['formula1'] == dv['source'] and v['sqref'] == dv['sqref']) for v in validations))
                    self.assertIn(f['target']['anchor'], main.comments)
                    lookup = wb.sheet('Lookup Tables')
                    self.assertEqual([lookup.value('L%d' % i) for i in range(1, 8)], ['SO%d' % i for i in range(1, 8)])

    def test_derived_formula_cells_are_vlookups(self):
        for t in MAPPING['templates']:
            with open_template(t['template_id']) as wb:
                for row in t['so_rows']:
                    ref = 'B%d' % row['row']
                    self.assertIn("VLOOKUP(A%d,'Lookup Tables'!$L$1:$M$7" % row['row'], wb.sheet('Main').formula(ref))

    def test_every_nonempty_main_cell_is_accounted_for(self):
        for t in MAPPING['templates']:
            with self.subTest(template=t['template_id']), open_template(t['template_id']) as wb:
                main = wb.sheet('Main')
                covered = set()
                def add_range(rng):
                    c1, r1, c2, r2 = split_range(rng)
                    covered.update(num_to_col(c) + str(r) for r in range(r1, r2 + 1) for c in range(c1, c2 + 1))
                for f in self.all_fields(t):
                    add_range(f['target']['range'])
                    if f['label']:
                        add_range(f['label']['cell'])
                for item in t['static_cells']:
                    add_range(item['cell'])
                    self.assertEqual(main.value(item['cell']), item['text'])
                for item in t['assessment_block']['column_headings'].values():
                    add_range(item['cell'])
                    self.assertEqual(main.value(item['cell']), item['text'])
                for rule in t['cleared_rows']:
                    for c in rule['cells_with_template_content']:
                        add_range(c['range'])
                for item in t['unresolved_regions']:
                    add_range(item['range'])
                    self.assertIn(item['range'], main.merged)
                    self.assertFalse(main.range_values(item['range']))
                self.assertEqual(set(main.cells) - covered, set())

    def test_cip_sources_resolve(self):
        files = set(MAPPING['cip_sources'].values())
        for name in files:
            self.assertTrue((ROOT / name).is_file(), name)
        assessment = load('ee-assessment.json')
        self.assertIn('active_term', assessment)
        self.assertIn('so_course_selection', assessment)
        course = next(c for c in CURRICULUM['curriculum']['courses'] if c['course_code'] == 'EE 322')
        for attr in ('course_title', 'course_code', 'instructor', 'clos'):
            self.assertIn(attr, course)
        for clo in course['clos']:
            for attr in ('clo_number', 'clo_text', 'mapped_sos', 'pi_codes', 'assessment_methods'):
                self.assertIn(attr, clo)
        for t in MAPPING['templates']:
            for f in self.all_fields(t):
                if f['cip_source']:
                    self.assertIn(f['cip_source']['file'], files)

    def test_multi_assessment_limitation_recorded(self):
        for t in MAPPING['templates']:
            block = t['assessment_block']
            self.assertEqual(block['row'], 10)
            self.assertEqual([f['target']['range'] for f in block['fields']], ['A10:B10', 'C10:E10', 'F10:J10'])
            self.assertIn('QAC-13', block['fields'][0]['flags'])


class ReviewIssuesTests(unittest.TestCase):
    def test_evidence_matches_templates(self):
        ids = [i['id'] for i in ISSUES['issues']]
        self.assertEqual(len(ids), len(set(ids)))
        known = {t['template_id'] for t in TEMPLATES['templates']}
        for issue in ISSUES['issues']:
            self.assertTrue(set(issue['templates']) <= known)
            self.assertIn(issue['status'], ISSUES['status_values'])
            self.assertEqual(issue['resolution'] is None, issue['status'] == 'open')
            for ev in issue['evidence']:
                with self.subTest(issue=issue['id'], cell=ev['cell']), open_template(ev['template_id']) as wb:
                    sheet = wb.sheet(ev['sheet'])
                    value, formula = sheet.value(ev['cell']), sheet.formula(ev['cell'])
                    self.assertTrue(value is not None or formula is not None, 'evidence cell is empty')
                    if ev['contains']:
                        self.assertIn(ev['contains'], '%s\n%s' % (value, formula))


class ApprovedDataUnchangedTests(unittest.TestCase):
    """Phase 1 must not alter approved curriculum data, QAC templates' source, or portal files."""

    def setUp(self):
        if git('rev-parse', '--verify', BASELINE_COMMIT + '^{commit}').returncode != 0:
            self.skipTest('baseline commit %s not available' % BASELINE_COMMIT)

    def test_only_expected_tracked_files_modified(self):
        result = git('diff', '--name-status', BASELINE_COMMIT)
        self.assertEqual(result.returncode, 0, result.stderr)
        changed = {}
        for line in result.stdout.splitlines():
            status, path = line.split('\t', 1)
            changed[path] = status
        offending = {p: s for p, s in changed.items()
                     if s[0] in 'MDRT' and p not in PHASE1_MAY_MODIFY | {CURRICULUM_PATH, EVIDENCE_PATH}}
        self.assertEqual(offending, {}, 'existing files changed outside the allowed Phase 1 set')

    def test_approved_curriculum_files_unchanged_since_baseline(self):
        tree = git('ls-tree', '-r', '--name-only', BASELINE_COMMIT, '--', 'data').stdout.splitlines()
        self.assertIn('data/ee_curriculum.json', tree)
        for path in tree:
            if path in PHASE1_MAY_MODIFY or path in (CURRICULUM_PATH, EVIDENCE_PATH):
                continue
            with self.subTest(path=path):
                self.assertEqual(git('diff', '--quiet', BASELINE_COMMIT, '--', path).returncode, 0, path + ' differs from baseline')

    def test_curriculum_differs_from_baseline_only_by_pi31_statement(self):
        baseline = json.loads(git('show', '%s:%s' % (BASELINE_COMMIT, CURRICULUM_PATH)).stdout)
        pis = baseline['abet']['performance_indicators']
        self.assertEqual(pis['PI31']['statement'], PI31_ORIGINAL)
        pis['PI31']['statement'] = PI31_ADOPTED
        self.assertEqual(CURRICULUM, baseline, 'ee_curriculum.json may differ from baseline only in the PI31 statement')

    def test_alignment_evidence_differs_from_baseline_only_in_pi31_profile(self):
        baseline = json.loads(git('show', '%s:%s' % (BASELINE_COMMIT, EVIDENCE_PATH)).stdout)
        current = load('ee_alignment_evidence.json')
        old, new = baseline['pi_evidence_profiles']['PI31'], current['pi_evidence_profiles']['PI31']
        self.assertNotEqual(old, new)
        self.assertEqual(set(old), set(new))
        changed = {k for k in old if old[k] != new[k]}
        self.assertEqual(changed, {'demonstrates', 'relationship', 'rubric_criteria', 'attainment_methodology'})
        current['pi_evidence_profiles']['PI31'] = old
        self.assertEqual(current, baseline, 'ee_alignment_evidence.json may differ only in the PI31 profile text')

    def test_runtime_allowlist_does_not_expose_qac_assets(self):
        allowlist = json.loads((ROOT / 'packaging' / 'desktop-runtime-files.json').read_text(encoding='utf-8'))
        self.assertFalse([p for p in allowlist if 'qac' in p.lower()])


if __name__ == '__main__':
    unittest.main()
